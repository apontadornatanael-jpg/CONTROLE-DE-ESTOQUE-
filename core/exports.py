"""Exportação de relatórios em Excel e PDF (sem dependência do Streamlit)."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AZUL = "#18324a"


# ----------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------
def _sem_timezone(df: pd.DataFrame) -> pd.DataFrame:
    """O Excel não aceita datas com fuso horário."""
    df = df.copy()
    for col in df.columns:
        if isinstance(df[col].dtype, pd.DatetimeTZDtype):
            df[col] = df[col].dt.tz_localize(None)
    return df


def excel_bytes(abas: dict[str, pd.DataFrame]) -> bytes:
    """Gera um .xlsx com cabeçalho estilizado, filtro e colunas ajustadas."""
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for nome, df in abas.items():
            df = _sem_timezone(df)
            nome_aba = nome[:31]
            df.to_excel(writer, sheet_name=nome_aba, index=False)
            ws = writer.sheets[nome_aba]
            for cel in ws[1]:
                cel.font = Font(bold=True, color="FFFFFF")
                cel.fill = PatternFill("solid", fgColor=AZUL.lstrip("#"))
                cel.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            if df.shape[1]:
                ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(df.columns, start=1):
                # str(v) em cada valor: células vazias (None/NaN) não podem quebrar
                amostra = df[col].head(200).map(lambda v: len(str(v)) if pd.notna(v) else 0)
                largura = max(len(str(col)), int(amostra.max()) if len(amostra) else 0)
                ws.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 50)
    return buf.getvalue()


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
def _estilos():
    base = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle("t", parent=base["Title"], textColor=colors.HexColor(AZUL), fontSize=18, alignment=0),
        "sub": ParagraphStyle("s", parent=base["Normal"], textColor=colors.HexColor("#6b7280"), fontSize=9),
        "normal": ParagraphStyle("n", parent=base["Normal"], fontSize=10, leading=13),
        "celula": ParagraphStyle("c", parent=base["Normal"], fontSize=9, leading=11),
    }


def _p(texto, estilo) -> Paragraph:
    return Paragraph(escape(str(texto if texto is not None else "")), estilo)


def _tabela(dados, larguras, destaque_linhas: dict | None = None) -> Table:
    t = Table(dados, colWidths=larguras, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(AZUL)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d5db")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f8")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for linha, cor in (destaque_linhas or {}).items():
        estilo.append(("BACKGROUND", (0, linha), (-1, linha), colors.HexColor(cor)))
    t.setStyle(TableStyle(estilo))
    return t


def comprovante_pdf(reg: dict) -> bytes:
    """Comprovante de uma movimentação.

    reg: numero, tipo, data_hora(datetime), destino, itens[{descricao,
    quantidade, unidade, numero_serie}], finalidade, responsavel, aprovador,
    emergencial, divergencia.
    """
    est = _estilos()
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=1.8 * cm, bottomMargin=1.8 * cm,
        title=f"Comprovante {reg['numero']}",
    )
    eh_retirada = reg["tipo"] == "RETIRADA"
    titulo = "Comprovante de retirada de material" if eh_retirada else "Comprovante de entrada de material"
    el = [
        _p(titulo, est["titulo"]),
        _p(f"Movimentação {reg['numero']} · {reg['data_hora'].strftime('%d/%m/%Y %H:%M')}", est["sub"]),
        Spacer(1, 0.5 * cm),
    ]
    meta = []
    if reg.get("destino"):
        meta.append(("Destino", reg["destino"]))
    meta.append(("Finalidade" if eh_retirada else "Origem / observação", reg.get("finalidade") or "—"))
    meta.append(("Retirado por" if eh_retirada else "Recebido por", reg.get("responsavel") or "—"))
    if reg.get("aprovador"):
        meta.append(("Aprovado por", reg["aprovador"]))
    if reg.get("emergencial"):
        meta.append(("Atenção", "RETIRADA EMERGENCIAL"))
    if reg.get("divergencia"):
        meta.append(("Divergência", reg["divergencia"]))
    for rotulo, valor in meta:
        el.append(Paragraph(f"<b>{escape(rotulo)}:</b> {escape(str(valor))}", est["normal"]))
    el.append(Spacer(1, 0.5 * cm))

    dados = [["#", "Material", "Quantidade", "Unidade", "Nº série"]]
    for i, it in enumerate(reg["itens"], start=1):
        dados.append(
            [
                str(i),
                _p(it["descricao"], est["celula"]),
                f"{it['quantidade']:g}",
                it["unidade"],
                _p(it.get("numero_serie") or "—", est["celula"]),
            ]
        )
    el.append(_tabela(dados, [1 * cm, 8 * cm, 2.6 * cm, 2.2 * cm, 3.2 * cm]))
    el.append(Spacer(1, 2.2 * cm))

    assin = [[
        "______________________________\nResponsável",
        "______________________________\nAprovador" if eh_retirada else "______________________________\nConferente",
    ]]
    t = Table(assin, colWidths=[8.5 * cm, 8.5 * cm])
    t.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 9)]))
    el.append(t)
    doc.build(el)
    return buf.getvalue()


def relatorio_estoque_pdf(df: pd.DataFrame, gerado_em: datetime) -> bytes:
    """Posição de estoque; linhas zeradas/baixas ganham destaque."""
    est = _estilos()
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, title="Posição de estoque",
    )
    el = [
        _p("Posição de estoque", est["titulo"]),
        _p(f"Gerado em {gerado_em.strftime('%d/%m/%Y %H:%M')} · {len(df)} materiais", est["sub"]),
        Spacer(1, 0.4 * cm),
    ]
    dados = [["Código", "Descrição", "Unid.", "Saldo", "Mínimo", "Máximo", "Situação"]]
    destaque = {}
    for i, r in enumerate(df.itertuples(index=False), start=1):
        dados.append(
            [
                _p(r[0], est["celula"]), _p(r[1], est["celula"]), r[2],
                f"{r[3]:g}", f"{r[4]:g}", f"{r[5]:g}", r[6],
            ]
        )
        if r[6] == "Zerado":
            destaque[i] = "#fee2e2"
        elif r[6] == "Baixo":
            destaque[i] = "#fef3c7"
    el.append(_tabela(dados, [3 * cm, 10 * cm, 2 * cm, 2.6 * cm, 2.6 * cm, 2.6 * cm, 2.8 * cm], destaque))
    doc.build(el)
    return buf.getvalue()
